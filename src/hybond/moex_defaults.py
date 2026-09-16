"""Import MOEX's emidocs-default disclosure (obligation-level default/
technical-default events across the whole exchange, not just the ВДО
watchlist) into moex_default, and serve lookups against it.

Source: <date>-emidocs-default.xlsx, sheet "emidocs-default" -- one row per
(issuer, ISIN, obligation type, planned date) with the disclosed status
('Дефолт' confirmed, or 'Технический дефолт' -- a missed-then-cured or
still-open short delay). Joins to hybond's own issuer table happen by ИНН
(inn = issuer.issuer_id), the same key convention used throughout this
codebase (see pd_import.py, universe_pipeline.py).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from .store import db

SHEET_NAME = "emidocs-default"
EXPECTED_HEADER = [
    "Краткое наименование эмитента", "ИНН", "ISIN",
    "Вид обязательства", "Состояние", "Плановая дата",
]
CONFIRMED_STATUS = "Дефолт"


def _iso_date(raw) -> str | None:
    """openpyxl already returns a datetime for this column -- ISO-format it
    for consistent, sortable storage. Accepts a plain string as a fallback
    (e.g. a hand-edited copy of the source file) rather than raising."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if hasattr(raw, "isoformat"):
        return raw.isoformat()
    return str(raw).strip() or None


def import_moex_defaults(xlsx_path: str, con) -> dict:
    """Import the emidocs-default sheet into moex_default.

    Idempotent: re-running against the same file (or a refreshed file with
    overlapping rows) leaves already-imported rows untouched via INSERT OR
    IGNORE against the (isin, obligation_type, planned_date) unique key.

    Returns: {rows_imported, rows_skipped, watchlist_matches} -- the last a
    list of {inn, issuer_name, isin, status, planned_date} dicts for rows
    whose ИНН matches an issuer already in hybond's own issuer table (any
    tier, active or not), computed AFTER the import so a fresh file's new
    rows are included in the match set.

    Raises: FileNotFoundError if xlsx_path doesn't exist; ValueError if the
    emidocs-default sheet is missing or its header doesn't match the
    expected columns.
    """
    import openpyxl  # lazy: only this function needs it

    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"MOEX defaults source not found: {xlsx_path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Sheet {SHEET_NAME!r} not found in {xlsx_path} "
            f"(sheets present: {wb.sheetnames})")
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError(f"Sheet {SHEET_NAME!r} in {xlsx_path} is empty")
    header = [str(h).strip() if h is not None else h for h in header]
    if header != EXPECTED_HEADER:
        raise ValueError(
            f"Sheet {SHEET_NAME!r} has unexpected columns {header}, "
            f"expected {EXPECTED_HEADER}")

    db.init_moex_default(con)

    rows_imported = 0
    rows_skipped = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        issuer_name, inn, isin, obligation_type, status, planned_date_raw = row[:6]
        planned_date = _iso_date(planned_date_raw)
        if isin is None and planned_date is None:
            rows_skipped += 1
            continue
        inn = str(inn).strip() if inn is not None else None
        isin = str(isin).strip() if isin is not None else None

        cur = con.execute(
            "INSERT OR IGNORE INTO moex_default "
            "(issuer_name, inn, isin, obligation_type, status, planned_date) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (str(issuer_name).strip() if issuer_name is not None else None,
             inn, isin, obligation_type, status, planned_date))
        if cur.rowcount:
            rows_imported += 1
        else:
            rows_skipped += 1

    con.commit()
    return {
        "rows_imported": rows_imported,
        "rows_skipped": rows_skipped,
        "watchlist_matches": get_watchlist_defaults(con),
    }


def get_watchlist_defaults(con) -> list[dict]:
    """All moex_default rows whose inn matches an issuer in the issuer
    table (any tier, active or not) -- joined on inn = issuer_id. Ordered
    by issuer name, then planned_date, for a stable display order.

    The join still spans both tiers, but the tier itself is deliberately
    NOT selected: watchlist_tier is an internal pipeline scoping field and
    was retired from everything the dashboard is handed (Купон Лаб redesign
    stage 3, 2026-09-16). See AGENTS.md."""
    db.init_moex_default(con)
    rows = con.execute(
        "SELECT d.inn, i.name AS issuer_name, d.isin, "
        "       d.obligation_type, d.status, d.planned_date "
        "FROM moex_default d JOIN issuer i ON i.issuer_id = d.inn "
        "ORDER BY i.name, d.planned_date").fetchall()
    return [dict(r) for r in rows]


def get_defaults_for_issuer(con, issuer_id: str) -> list[dict]:
    """All moex_default rows for one issuer_id (inn) -- used by the bond
    drawer to show default history. Ordered oldest first."""
    db.init_moex_default(con)
    rows = con.execute(
        "SELECT issuer_name, inn, isin, obligation_type, status, planned_date "
        "FROM moex_default WHERE inn = ? ORDER BY planned_date",
        (issuer_id,)).fetchall()
    return [dict(r) for r in rows]
