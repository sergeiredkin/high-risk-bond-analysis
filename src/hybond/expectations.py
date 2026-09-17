"""What the CBR and the analysts it surveys expect for the key rate -- with
history, publication by publication (Ставка и ОФЗ page, plan rev 2 approved
2026-09-17).

Sources, each read from the official page it is updated on (every URL below
was opened and its content checked on 2026-09-17):

  cbr_forecast  Среднесрочный прогноз, one PDF per anchor meeting, listed on
                https://www.cbr.ru/dkp/mp_dec/decision_key_rate/ -- average key
                rate by year plus the «остаток года» footnote. 11 PDFs back to
                16.02.2024 on that page.
  ondkp         Основные направления ЕГДКП, https://www.cbr.ru/about_br/publ/ondkp/
                -- the four scenarios (Базовый, Дезинфляционный, Проинфляционный,
                Рисковый) as HTML tables. Published twice a year on ONE URL: a
                draft ("Проект от 31 августа 2026 года") later replaced by the
                final edition, so each version is stored separately.
  survey        Макроэкономический опрос, https://www.cbr.ru/statistics/ddkp/mo_br/
                -- the whole history in one workbook (full.xlsx, key-rate sheet):
                median, mean, quartiles, 10/90th percentiles, min/max and number
                of answers for every survey since May 2021.
  documents     press releases, Chairman's statements, forecasts, discussion
                summaries and comments, indexed from the decisions page.
  key_rate      the daily key rate, https://www.cbr.ru/hd_base/KeyRate/ (from 2013).
  calendar      https://www.cbr.ru/dkp/cal_mp/, via cbr_tracker.refresh_schedule().

The CBR has no structured feed of media interviews (/press/interview/ and
/press/speeches/ return 404); by the captain's decision (Q2) they are not
collected.

Every source refreshes independently: one failing page logs a pipeline_health
warning and leaves the others -- and its own last good data -- untouched.
"""
from __future__ import annotations

import html as html_mod
import io
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import requests

from . import cbr_tracker
from .cbr_tracker import (
    _MONTHS_GEN,
    _get,
    _html_to_text,
    _pdf_to_text,
    forecast_rate_by_year,
    forecast_rest_of_year,
    parse_rate_range,
)
from .health import log_health
from .store import db

WWW = "https://www.cbr.ru"
DECISIONS_URL = f"{WWW}/dkp/mp_dec/decision_key_rate/"
ONDKP_INDEX_URL = f"{WWW}/about_br/publ/ondkp/"
SURVEY_PAGE_URL = f"{WWW}/statistics/ddkp/mo_br/"
KEY_RATE_URL = f"{WWW}/hd_base/KeyRate/"
CALENDAR_URL = f"{WWW}/dkp/cal_mp/"
RAW_DIR = Path("data/raw/cbr/expectations")

# Where the page links each source to. Kept in one place so the export, the
# CLI and the tests agree on the exact official URL.
SOURCE_PAGES = {
    "cbr_forecast": DECISIONS_URL,
    "ondkp": ONDKP_INDEX_URL,
    "survey": SURVEY_PAGE_URL,
    "documents": DECISIONS_URL,
    "key_rate": KEY_RATE_URL,
    "calendar": CALENDAR_URL,
}

SCENARIO_KEYS = {
    "Базовый": "base",
    "Дезинфляционный": "disinflation",
    "Проинфляционный": "proinflation",
    "Рисковый": "risk",
}

SURVEY_STATS = {
    "Медиана": "median", "Среднее": "mean", "Макс": "max", "90-й процентиль": "p90",
    "3-й квартиль": "q3", "1-й квартиль": "q1", "10-й процентиль": "p10", "Мин": "min",
    "Всего ответов": "n",
}

_MONTHS_NOM = {"январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5, "июнь": 6, "июль": 7,
               "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12}

# Politeness bounds for the daily run; --backfill lifts them.
DAILY_DOCUMENT_BODIES = 4


def _absolute(url: str) -> str:
    return url if url.startswith("http") else WWW + url


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", html_mod.unescape(s).replace("\xa0", " ")).strip()


def _ru_date(day: str, month_gen: str, year: str) -> str | None:
    """('31', 'августа', '2026') -> '2026-08-31'; None for an unknown month."""
    month = _MONTHS_GEN.get(month_gen.lower())
    return f"{int(year):04d}-{month:02d}-{int(day):02d}" if month else None


def _dmy(s: str) -> str:
    d, m, y = s.split(".")
    return f"{y}-{m}-{d}"


def _snapshot(name: str, content: bytes | str) -> None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    mode = "wb" if isinstance(content, bytes) else "w"
    with open(RAW_DIR / name, mode, encoding=None if mode == "wb" else "utf-8") as fh:
        fh.write(content)


# ---- parsers (pure) ----------------------------------------------------------

_NEWS_DATE_RE = re.compile(
    r'news-info-line_date">\s*(Проект от\s+)?(\d{1,2})(?:&nbsp;|\s)+([а-я]+)(?:&nbsp;|\s)+(\d{4})')


def parse_page_date(page_html: str) -> tuple[str | None, bool]:
    """The date line cbr.ru prints under a publication's title ->
    (ISO date, is_draft). ОНЕГДКП drafts read "Проект от 31 августа 2026
    года"; the final edition and the survey page print the bare date."""
    m = _NEWS_DATE_RE.search(page_html)
    if not m:
        return None, False
    return _ru_date(m.group(2), m.group(3), m.group(4)), bool(m.group(1))


def parse_ondkp_index(page_html: str) -> list[str]:
    """Edition slugs listed on the ОНЕГДКП index, newest first:
    ['on_2027_2029', 'on_2026_2028', ...]."""
    slugs = sorted(set(re.findall(r"/about_br/publ/ondkp/(on_\d{4}_\d{4})/", page_html)), reverse=True)
    return slugs


_SCENARIO_HEAD_RE = re.compile(r"(Базовый сценарий|Сценарий\s*(?:&#171;|«)\s*([А-Яа-я]+)\s*(?:&#187;|»))")


def parse_ondkp_edition(page_html: str) -> dict[str, Any]:
    """One ОНЕГДКП edition page -> {published, status, scenarios}.

    scenarios = {'base'|'disinflation'|'proinflation'|'risk': {year: (lo, hi)}},
    read from the «Ключевая ставка, в среднем за год» row of the table that
    follows each scenario heading; the prior-year actual column is dropped.
    status is 'draft' for "Проект от ..." and 'final' otherwise."""
    published, is_draft = parse_page_date(page_html)
    scenarios: dict[str, dict[int, tuple[float, float]]] = {}
    for tm in re.finditer(r"<table.*?</table>", page_html, re.DOTALL):
        heads = list(_SCENARIO_HEAD_RE.finditer(page_html, 0, tm.start()))
        if not heads:
            continue
        head = heads[-1]
        name = "Базовый" if head.group(1).startswith("Базовый") else head.group(2)
        key = SCENARIO_KEYS.get(name)
        if not key or key in scenarios:
            continue
        table = tm.group(0)
        years = [int(y) for y in re.findall(r"<th[^>]*>\s*(20\d\d)", table)]
        rate_row = next((r for r in re.findall(r"<tr.*?</tr>", table, re.DOTALL)
                         if "Ключевая ставка" in r), None)
        if not years or not rate_row:
            continue
        cells = [_clean(re.sub(r"<[^>]+>", " ", c)) for c in re.findall(r"<td[^>]*>(.*?)</td>", rate_row, re.DOTALL)]
        values = cells[1:]  # first cell is the row label
        if len(values) != len(years):
            continue
        by_year = {}
        for y, v in zip(years[1:], values[1:]):  # drop the факт column
            rng = parse_rate_range(v)
            if rng:
                by_year[y] = rng
        if by_year:
            scenarios[key] = by_year
    return {"published": published, "status": "draft" if is_draft else "final", "scenarios": scenarios}


_SURVEY_PERIOD_RE = re.compile(r"Результаты опроса:\s*([а-я]+)\s+(\d{4})\s+года")


def parse_survey_page(page_html: str) -> dict[str, Any]:
    """The macro-survey page -> {published, period ('2026-09'), xlsx_url}."""
    published, _ = parse_page_date(page_html)
    m = _SURVEY_PERIOD_RE.search(_clean(page_html))
    period = None
    if m and m.group(1).lower() in _MONTHS_NOM:
        period = f"{int(m.group(2)):04d}-{_MONTHS_NOM[m.group(1).lower()]:02d}"
    x = re.search(r'href="([^"]*full\.xlsx)"', page_html)
    return {"published": published, "period": period, "xlsx_url": _absolute(x.group(1)) if x else None}


def parse_survey_workbook(content: bytes) -> list[dict[str, Any]]:
    """full.xlsx -> [{period: 'YYYY-MM', year, stat, value}] for the key-rate
    sheet (average key rate for the year, including weekends).

    The sheet is found through the workbook's own navigation sheet rather
    than by a fixed name; '-' cells (no answer for that year) are skipped."""
    import openpyxl  # lazy: only the survey refresh needs it

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet_name = None
    for ws in wb.worksheets:
        if not ws.title.lower().startswith("навигация"):
            continue
        for row in ws.iter_rows(values_only=True):
            label = next((c for c in row if isinstance(c, str) and c.strip()), "")
            num = next((c for c in row if isinstance(c, (int, float))), None)
            if label.strip().startswith("Ключевая ставка") and num is not None:
                sheet_name = str(int(num))
                break
    if sheet_name is None or sheet_name not in wb.sheetnames:
        raise ValueError("survey workbook: key-rate sheet not found in the navigation sheet")

    rows = list(wb[sheet_name].iter_rows(values_only=True))
    header = next((r for r in rows if any(isinstance(c, str) and c.strip() == "Прогнозный период" for c in r)), None)
    if header is None:
        raise ValueError("survey workbook: 'Прогнозный период' header row not found")
    period_cols = {i: c for i, c in enumerate(header) if isinstance(c, datetime)}
    out: list[dict[str, Any]] = []
    stat = None
    for r in rows:
        label = r[1] if len(r) > 1 and isinstance(r[1], str) else None
        if label and label.strip() in SURVEY_STATS:
            stat = SURVEY_STATS[label.strip()]
        year_cell = next((c for c in r[:4] if isinstance(c, datetime)), None)
        if stat is None or year_cell is None:
            continue
        for i, period in period_cols.items():
            v = r[i] if i < len(r) else None
            if isinstance(v, (int, float)):
                out.append({"period": f"{period.year:04d}-{period.month:02d}", "year": year_cell.year,
                            "stat": stat, "value": float(v)})
    if not out:
        raise ValueError("survey workbook: key-rate sheet parsed to no values")
    return out


_DOC_ENTRY_RE = re.compile(
    r'document-regular_date">\s*([\d.]+)\s*</div>.*?href="([^"]+)"[^>]*>\s*'
    r'<span class="document-regular_name_visible">(.*?)</span>', re.DOTALL)


def article_text(page_html: str) -> str:
    """Visible text of a cbr.ru article page, title to the page's own
    «Страница была полезной?» footer. cbr_tracker._html_to_text() keeps only
    the landing-text blocks, which on a discussion-summary page stops before
    the decision options -- this keeps the whole article."""
    body = re.sub(r"<(script|style)\b[^>]*>.*?</\1>", " ", page_html, flags=re.DOTALL | re.IGNORECASE)
    start = body.find("<h1")
    body = body[start if start != -1 else 0:]
    text = _clean(re.sub(r"<[^>]+>", " ", body))
    end = text.find("Страница была полезной")
    return text[:end].strip() if end != -1 else text


def _doc_type(url: str) -> str | None:
    url = url.lower()
    if "/press/pr/" in url and "key" in url:
        return "press_release"
    if "/press/event/" in url:
        return "statement"
    if re.search(r"forecast_\d+\.pdf", url):
        return "forecast"
    if "summary_key_rate_" in url:
        return "summary"
    if "comment_" in url:
        return "comment"
    return None


def parse_decision_documents(page_html: str) -> list[dict[str, Any]]:
    """The decisions page -> [{meeting_date, doc_type, published, title, url}],
    one per document listed under each «Заседание Совета директоров от
    DD.MM.YYYY» block."""
    parts = re.split(r"Заседание Совета директоров от (\d{2}\.\d{2}\.\d{4})\s*</a>", page_html)
    out = []
    for i in range(1, len(parts) - 1, 2):
        meeting = _dmy(parts[i])
        for published, url, title in _DOC_ENTRY_RE.findall(parts[i + 1]):
            kind = _doc_type(url)
            if not kind:
                continue
            out.append({"meeting_date": meeting, "doc_type": kind, "published": _dmy(published),
                        "title": _clean(re.sub(r"<[^>]+>", " ", title)), "url": _absolute(url)})
    return out


def parse_key_rate_table(page_html: str) -> list[tuple[str, float]]:
    """hd_base/KeyRate -> [(ISO date, percent)], oldest first."""
    rows = re.findall(r"<td>\s*(\d{2}\.\d{2}\.\d{4})\s*</td>\s*<td>\s*(\d+,\d+)\s*</td>", page_html)
    return sorted((_dmy(d), float(v.replace(",", "."))) for d, v in rows)


def extract_signal(press_text: str | None) -> str | None:
    """The forward-guidance sentence of a key-rate press release: the first
    sentence that starts «Банк России» and says what it will do next
    («будет принимать дальнейшие решения…», «будет оценивать
    целесообразность…», «допускает возможность повышения…»)."""
    if not press_text:
        return None
    for s in re.split(r"(?<=[.!?])\s+", press_text.replace("\n", " ")):
        if (s.startswith("Банк России") and "ставк" in s
                and re.search(r"дальнейш|целесообразн|ближайш|допуска|не исключа|вернуться|пауз", s)):
            return s.strip()
    return None


_DECISION_SECTION = "Выводы для денежно-кредитной политики и решение по ключевой ставке"
_OPTIONS_RE = re.compile(
    r"вариант\w*\s+решени\w*[^:.]{0,60}:\s*((?:(?:сохранени|снижени|повышени)\w*[^;]*?(?:годовых|пункт\w*)\s*[;.]\s*)+)",
    re.IGNORECASE)


def extract_summary_options(summary_text: str | None) -> dict[str, str | None]:
    """A «Резюме обсуждения ключевой ставки» page -> {options, positions}.

    options   -- the decision options the board discussed, joined with "; "
                 («сохранение ключевой ставки на уровне 14,25% годовых;
                 снижение … до 14,00% годовых»). None when the summary does
                 not list them (several 2024 summaries do not).
    positions -- up to two sentences, after the options, in which the
                 participants' view is stated («Основная часть участников
                 обсуждения отметили…»); argument-list headers are skipped.

    Both are read from the decision section at the end of the page, not the
    table of contents that repeats its title near the top."""
    if not summary_text:
        return {"options": None, "positions": None}
    flat = re.sub(r"\s+", " ", summary_text)
    at = flat.rfind(_DECISION_SECTION)
    section = flat[at + len(_DECISION_SECTION):] if at != -1 else flat
    m = _OPTIONS_RE.search(section)
    options = None
    rest = section
    if m:
        items = [i.strip() for i in re.split(r"\s*[;.]\s+(?=[А-ЯЁа-яё])|\s*[;.]\s*$", m.group(1)) if i.strip()]
        options = "; ".join(i[0].lower() + i[1:] for i in items)
        rest = section[m.end():]
    sentences = re.split(r"(?<=[а-яё»)%])\s*\.\s+(?=[А-ЯЁ])", rest)
    # argument-list headers ("Основные аргументы участников…", "…следующие:")
    # introduce bullet lists, they are not the participants' conclusion
    picked = [x.strip() for x in sentences
              if "участник" in x.lower() and not re.search(r"аргумент|следующ|рассмотрели", x)][:2]
    positions = ". ".join(picked) + "." if picked else None
    if positions and len(positions) > 600:
        positions = positions[:597].rstrip() + "…"
    return {"options": options, "positions": positions}


# ---- refreshers (network + store) ----------------------------------------------

def store_forecast_text(con, meeting_date: str, published: str, url: str, text: str) -> bool:
    """Parse one Среднесрочный прогноз PDF's text into rate_expectation
    (source cbr_forecast, edition = meeting date): the average key rate for
    every year in its table plus the «остаток года» footnote. Logs
    EXPECT_PARSE_FAIL and stores nothing when the key-rate row does not line
    up with the header years. Returns whether anything was stored."""
    by_year = forecast_rate_by_year(text)
    if not by_year:
        log_health(con, "cbr", "ingest", "warn", "EXPECT_PARSE_FAIL",
                   f"forecast {url}: key-rate row did not line up with the header years")
        return False
    today = date.today().isoformat()
    rows = [{"scenario": "base", "year": y, "lo": lo, "hi": hi, "url": url, "fetched_on": today}
            for y, (lo, hi) in by_year.items()]
    rest = forecast_rest_of_year(text)
    if rest:
        rows.append({"scenario": "base", "period": "rest_of_year", "year": rest["year"],
                     "lo": rest["lo"], "hi": rest["hi"], "ytd_avg": rest["ytd_avg"],
                     "rest_from": rest["rest_from"], "url": url, "fetched_on": today})
    db.init_rates(con)
    db.replace_rate_expectations(con, "cbr_forecast", meeting_date, published, rows)
    db.upsert_cbr_document(con, meeting_date=meeting_date, doc_type="forecast", url=url,
                           published=published, fetched_on=today)
    db.mark_expectation_source(con, "cbr_forecast", page_url=DECISIONS_URL, last_published=published)
    return True


def refresh_key_rate(con, full: bool = False) -> int:
    """Daily key rate: the whole series from 2013 when `full` or the store
    is empty, otherwise the last 30 days."""
    latest = db.latest_key_rate(con)
    start = date(2013, 1, 1) if full or not latest else date.fromisoformat(latest[0]) - timedelta(days=30)
    r = _get(KEY_RATE_URL, params={"UniDbQuery.Posted": "True",
                                  "UniDbQuery.From": start.strftime("%d.%m.%Y"),
                                  "UniDbQuery.To": date.today().strftime("%d.%m.%Y")})
    rows = parse_key_rate_table(r.text)
    if not rows:
        raise ValueError("key-rate table parse produced no rows; layout changed?")
    n = db.upsert_key_rates(con, rows)
    db.mark_expectation_source(con, "key_rate", page_url=KEY_RATE_URL, last_published=rows[-1][0])
    return n


def refresh_survey(con, force: bool = False) -> int:
    """Macro survey: re-download the workbook only when the page announces a
    survey the store does not have yet (or `force`). Stores every survey in
    the workbook -- the history is the point."""
    page = _get(SURVEY_PAGE_URL).text
    info = parse_survey_page(page)
    if not info["xlsx_url"] or not info["period"]:
        raise ValueError("survey page: workbook link or survey period not found; layout changed?")
    stored = {e for e, _, _ in db.rate_expectation_versions(con, "survey")}
    if info["period"] in stored and not force:
        db.mark_expectation_source(con, "survey", page_url=SURVEY_PAGE_URL, last_published=info["published"])
        return 0
    content = _get(info["xlsx_url"]).content
    _snapshot("survey_full.xlsx", content)
    values = parse_survey_workbook(content)
    today = date.today().isoformat()
    by_period: dict[str, list[dict]] = {}
    for v in values:
        by_period.setdefault(v["period"], []).append(v)
    n = 0
    for period, vals in by_period.items():
        rows = [{"scenario": v["stat"], "year": v["year"], "lo": v["value"], "hi": v["value"],
                 "url": info["xlsx_url"], "fetched_on": today} for v in vals]
        n += db.replace_rate_expectations(con, "survey", period, f"{period}-01", rows)
    db.mark_expectation_source(con, "survey", page_url=SURVEY_PAGE_URL, last_published=info["published"])
    return n


def refresh_ondkp(con, editions: int = 1) -> dict[str, Any]:
    """ОНЕГДКП: fetch the newest `editions` edition pages and store each
    version seen (a draft and its final edition live on the same URL, so a
    changed date or status is a new version, not an overwrite).

    The newest edition must parse into all four scenarios or this raises
    (a layout change must be loud). An older edition with a different
    scenario set -- 2025-2027 published only its base table in this form --
    is reported as skipped, not stored and not logged as a fault.
    Returns {rows, stored: [slug], skipped: [slug]}."""
    index = _get(ONDKP_INDEX_URL).text
    slugs = parse_ondkp_index(index)
    if not slugs:
        raise ValueError("ОНЕГДКП index: no editions found; layout changed?")
    today = date.today().isoformat()
    out: dict[str, Any] = {"rows": 0, "stored": [], "skipped": []}
    newest_published = None
    for i, slug in enumerate(slugs[:editions]):
        url = f"{ONDKP_INDEX_URL}{slug}/"
        parsed = parse_ondkp_edition(_get(url).text)
        if not parsed["published"] or len(parsed["scenarios"]) < 4:
            if i == 0:
                raise ValueError(f"ОНЕГДКП {slug}: expected 4 scenario tables, parsed "
                                 f"{sorted(parsed['scenarios'])}; layout changed?")
            out["skipped"].append(slug)
            continue
        rows = [{"scenario": key, "year": y, "lo": lo, "hi": hi, "url": url, "fetched_on": today}
                for key, by_year in parsed["scenarios"].items() for y, (lo, hi) in by_year.items()]
        out["rows"] += db.replace_rate_expectations(con, "ondkp", slug, parsed["published"], rows,
                                                    status=parsed["status"])
        out["stored"].append(slug)
        newest_published = max(newest_published or "", parsed["published"])
    db.mark_expectation_source(con, "ondkp", page_url=ONDKP_INDEX_URL, last_published=newest_published)
    return out


def refresh_documents(con, bodies: int | None = DAILY_DOCUMENT_BODIES) -> dict[str, int]:
    """Index every document on the decisions page into cbr_document, then
    fetch forecast PDFs that are not parsed yet (always all of them -- one
    per anchor meeting) and up to `bodies` press-release/summary texts not
    stored yet (None = all, for a backfill). Press-release text already held
    in cbr_meeting is copied, not re-fetched."""
    page = _get(DECISIONS_URL).text
    docs = parse_decision_documents(page)
    if not docs:
        raise ValueError("decisions page: no documents parsed; layout changed?")
    today = date.today().isoformat()
    for d in docs:
        db.upsert_cbr_document(con, **d)
    con.commit()

    have_forecasts = {r["url"] for r in db.rate_expectations(con, "cbr_forecast")}
    forecasts = 0
    for d in (d for d in docs if d["doc_type"] == "forecast" and d["url"] not in have_forecasts):
        text = _pdf_to_text(_get(d["url"]).content)
        if store_forecast_text(con, d["meeting_date"], d["published"], d["url"], text):
            forecasts += 1
        con.commit()

    fetched = 0
    stored = {(r["meeting_date"], r["doc_type"]): r for r in db.cbr_documents(con)}
    for d in docs:
        if d["doc_type"] not in ("press_release", "summary"):
            continue
        if stored.get((d["meeting_date"], d["doc_type"]), {}).get("text"):
            continue
        # A press release already held in cbr_meeting is copied, not refetched.
        # Not résumé text: cbr_meeting.resume_text comes from _html_to_text,
        # which stops before the options the page needs (see article_text).
        meeting = db.get_cbr_meeting(con, d["meeting_date"]) if d["doc_type"] == "press_release" else None
        if meeting and meeting["press_release_text"]:
            db.upsert_cbr_document(con, meeting_date=d["meeting_date"], doc_type=d["doc_type"],
                                   url=d["url"], text=meeting["press_release_text"], fetched_on=today)
            continue
        if bodies is not None and fetched >= bodies:
            continue
        page = _get(d["url"]).text
        text = _html_to_text(page) if d["doc_type"] == "press_release" else article_text(page)
        if text:
            db.upsert_cbr_document(con, meeting_date=d["meeting_date"], doc_type=d["doc_type"],
                                   url=d["url"], text=text, fetched_on=today)
            fetched += 1
    con.commit()
    newest = max(d["published"] for d in docs)
    db.mark_expectation_source(con, "documents", page_url=DECISIONS_URL, last_published=newest)
    forecast_versions = db.rate_expectation_versions(con, "cbr_forecast")
    db.mark_expectation_source(con, "cbr_forecast", page_url=DECISIONS_URL,
                               last_published=max((p for _, p, _ in forecast_versions), default=None))
    return {"documents": len(docs), "forecasts_parsed": forecasts, "bodies_fetched": fetched}


def refresh_calendar(con) -> int:
    """The CBR calendar into cbr_schedule (see cbr_tracker.refresh_schedule)."""
    n = cbr_tracker.refresh_schedule(con)
    upcoming = [m["date"] for m in cbr_tracker.get_schedule(con) if m["date"] >= date.today().isoformat()]
    db.mark_expectation_source(con, "calendar", page_url=CALENDAR_URL,
                               last_published=max((m["date"] for m in cbr_tracker.get_schedule(con)), default=None))
    if not upcoming:
        log_health(con, "cbr", "ingest", "warn", "CBR_SCHEDULE_ENDS",
                   "no future key-rate meeting in the CBR calendar yet; next year's dates not published")
    return n


def run_expectations(backfill: bool = False) -> dict[str, dict[str, Any]]:
    """Refresh every source; one failure never stops the others. Returns
    {source: {ok, result|error}}. `backfill` downloads full histories (key
    rate since 2013, the last three ОНЕГДКП editions, every document body);
    it is switched on automatically while no ОНЕГДКП edition is stored."""
    con = db.connect()
    db.init_db(con)
    db.init_rates(con)
    db.init_cbr(con)
    db.init_pipeline_health(con)
    # A store that has never held a scenario edition gets the whole history on
    # its first cron run -- the page is about how expectations moved, and a
    # first day with only today's numbers would not show that.
    if not backfill and not db.rate_expectation_versions(con, "ondkp"):
        backfill = True
    steps = [
        ("calendar", lambda: refresh_calendar(con)),
        ("key_rate", lambda: refresh_key_rate(con, full=backfill)),
        ("documents", lambda: refresh_documents(con, bodies=None if backfill else DAILY_DOCUMENT_BODIES)),
        ("ondkp", lambda: refresh_ondkp(con, editions=3 if backfill else 1)),
        ("survey", lambda: refresh_survey(con, force=backfill)),
    ]
    report: dict[str, dict[str, Any]] = {}
    for name, fn in steps:
        try:
            report[name] = {"ok": True, "result": fn()}
        except (requests.RequestException, ValueError, KeyError, OSError) as exc:
            con.rollback()
            detail = f"{type(exc).__name__}: {exc}"[:300]
            for source in (("documents", "cbr_forecast") if name == "documents" else (name,)):
                db.mark_expectation_source(con, source, page_url=SOURCE_PAGES.get(source), ok=False, error=detail)
            log_health(con, "cbr", "ingest", "warn", "EXPECT_FETCH_FAIL", f"{name}: {detail}")
            report[name] = {"ok": False, "error": detail}
        con.commit()
    con.close()
    return report
